from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

doc = Document()

# ── Page margins ──────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(2.8)
    section.right_margin  = Cm(2.8)

# ── Helper: set paragraph shading ────────────────────────────
def shade_paragraph(para, hex_color):
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    pPr.append(shd)

# ── Helper: add horizontal rule ──────────────────────────────
def add_hr(doc, color='E74C3C'):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    return p

# ════════════════════════════════════════════════════════════════
#  HEADER BANNER
# ════════════════════════════════════════════════════════════════
banner = doc.add_paragraph()
shade_paragraph(banner, '1A1E2A')
banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
banner.paragraph_format.space_before = Pt(6)
banner.paragraph_format.space_after  = Pt(6)
run = banner.add_run('  KALLYAS  —  Business Agency Website  ')
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
run.font.size = Pt(18)
run.font.bold = True
run.font.name = 'Calibri'

sub_banner = doc.add_paragraph()
shade_paragraph(sub_banner, 'E74C3C')
sub_banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_banner.paragraph_format.space_before = Pt(0)
sub_banner.paragraph_format.space_after  = Pt(12)
run2 = sub_banner.add_run('  Assignment 2 — Photoshop to HTML Conversion  ')
run2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
run2.font.size = Pt(11)
run2.font.name = 'Calibri'

# ════════════════════════════════════════════════════════════════
#  STUDENT INFO TABLE
# ════════════════════════════════════════════════════════════════
def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)
    r = p.add_run(text.upper())
    r.font.bold  = True
    r.font.size  = Pt(10)
    r.font.color.rgb = RGBColor(0xE7, 0x4C, 0x3C)
    r.font.name  = 'Calibri'
    add_hr(doc)

section_heading(doc, 'Student Information')

table = doc.add_table(rows=4, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.style = 'Table Grid'

info = [
    ('Student Name',   'Fazilma Khan Niazi'),
    ('Roll Number',    'S2024266184'),
    ('Assignment',     'Assignment 2 — Photoshop to HTML Conversion'),
    ('Submission Date', datetime.date.today().strftime('%B %d, %Y')),
]

for i, (label, value) in enumerate(info):
    row = table.rows[i]
    # Label cell
    lc = row.cells[0]
    lc.width = Inches(2)
    lp = lc.paragraphs[0]
    lr = lp.add_run(label)
    lr.font.bold = True
    lr.font.size = Pt(10)
    lr.font.name = 'Calibri'
    # shade label cells
    tc_pr = lc._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  'F2F2F2')
    tc_pr.append(shd)
    # Value cell
    vc = row.cells[1]
    vp = vc.paragraphs[0]
    vr = vp.add_run(value)
    vr.font.size = Pt(10)
    vr.font.name = 'Calibri'

doc.add_paragraph()

# ════════════════════════════════════════════════════════════════
#  PROJECT OVERVIEW
# ════════════════════════════════════════════════════════════════
section_heading(doc, 'Project Overview')

overview = doc.add_paragraph()
overview.paragraph_format.space_after = Pt(8)
r = overview.add_run(
    'This assignment required converting a Photoshop design template into a fully '
    'responsive static HTML webpage. The design template depicted a business agency '
    'website called "Kallyas". The final implementation closely matches the original '
    'design using Bootstrap 5 Grid for responsiveness and custom CSS for styling. '
    'No JavaScript was used — the website is built entirely with HTML and CSS.'
)
r.font.size = Pt(10)
r.font.name = 'Calibri'

# ════════════════════════════════════════════════════════════════
#  LIVE LINKS
# ════════════════════════════════════════════════════════════════
section_heading(doc, 'Live Links')

links = [
    ('Live Website (Netlify)', 'https://stirring-snickerdoodle-5a985f.netlify.app'),
    ('GitHub Repository',      'https://github.com/MianTaha0/coursera-test'),
    ('GitHub Branch',          'claude/beginner-assignment-2PP1B  →  assignment2/ folder'),
]

for label, url in links:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    bold_run = p.add_run(f'{label}: ')
    bold_run.font.bold = True
    bold_run.font.size = Pt(10)
    bold_run.font.name = 'Calibri'
    link_run = p.add_run(url)
    link_run.font.size = Pt(10)
    link_run.font.name = 'Calibri'
    link_run.font.color.rgb = RGBColor(0x0, 0x70, 0xC0)

doc.add_paragraph()

# ════════════════════════════════════════════════════════════════
#  WEBSITE SECTIONS
# ════════════════════════════════════════════════════════════════
section_heading(doc, 'Website Sections Implemented')

sections_list = [
    ('1. Navigation Bar',     'Fixed sticky navbar with brand logo and 6 nav links (Home, Services, Media, About Us, Platform, Contact). Responsive using CSS Flexbox.'),
    ('2. Hero Section',       'Full-screen background image with dark overlay, large heading, subtitle text, and two call-to-action buttons (Start Now, Our Services).'),
    ('3. Feature Cards',      'Three cards with icons — Content Concept & Strategy (red), Design & Concepts (dark), SEO & Marketing Solutions (darker). Overlapping hero bottom.'),
    ('4. Services Section',   'Six services in a 2-row Bootstrap grid: Strategy, Marketing, Technology, Ecommerce, Branding, SEO Identity. Each with icon and description.'),
    ('5. Dark CTA Section',   'Two-column dark background section with heading, description, button, and an image with a red accent decorative block.'),
    ('6. Why Us Section',     'Four numbered points (01–04) on the left in a 2×2 grid and descriptive text on the right explaining quality and cost-effectiveness.'),
    ('7. Portfolio Gallery',  'Eight portfolio images in a 2-row × 4-column grid with CSS hover overlay effect (red overlay with plus icon).'),
    ('8. Stats Section',      '"Trusted by 3600+ clients" heading with descriptive text, Learn More button, and team image with play button overlay.'),
    ('9. Clients Section',    'Dark background with four client logos/names (FastCompany, Southern Company, shield icon, V&IPS) and agency tagline.'),
    ('10. Features Section',  'Light gray background with heading, two feature boxes (Secured Database, Modern Framework with icons), and a person image.'),
    ('11. Contact Form',      'Two-column layout: heading on left, HTML form on right with name, email, phone, subject, message fields and submit button.'),
    ('12. Footer',            'Dark 4-column footer with brand info, address, social links, navigation links, platform links, and newsletter subscription form.'),
]

for title, desc in sections_list:
    p = doc.add_paragraph()
    p.paragraph_format.space_after  = Pt(5)
    p.paragraph_format.left_indent  = Cm(0.5)
    br = p.add_run(f'{title}: ')
    br.font.bold = True
    br.font.size = Pt(10)
    br.font.name = 'Calibri'
    dr = p.add_run(desc)
    dr.font.size = Pt(10)
    dr.font.name = 'Calibri'

# ════════════════════════════════════════════════════════════════
#  TECHNOLOGIES USED
# ════════════════════════════════════════════════════════════════
section_heading(doc, 'Technologies Used')

tech = [
    ('HTML5',           'Semantic structure — sections, nav, footer, form elements'),
    ('CSS3',            'Custom styles, CSS variables, Flexbox, transitions, hover effects, media queries'),
    ('Bootstrap 5',     'Responsive grid system (col-md-*, col-lg-*), utility classes'),
    ('Font Awesome 6',  'Icons for services, features, social links, and cards'),
    ('Google Fonts',    'Poppins font family (weights 300–800)'),
    ('Unsplash',        'High-quality placeholder images loaded via URL'),
    ('Netlify',         'Free static site hosting with GitHub integration'),
]

tech_table = doc.add_table(rows=len(tech) + 1, cols=2)
tech_table.style = 'Table Grid'
tech_table.alignment = WD_TABLE_ALIGNMENT.LEFT

# Header row
hdr_cells = tech_table.rows[0].cells
for cell, text in zip(hdr_cells, ['Technology', 'Usage']):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '1A1E2A')
    tc_pr.append(shd)
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r.font.size = Pt(10)
    r.font.name = 'Calibri'

for i, (tech_name, usage) in enumerate(tech):
    row = tech_table.rows[i + 1]
    r0 = row.cells[0].paragraphs[0].add_run(tech_name)
    r0.font.bold = True
    r0.font.size = Pt(10)
    r0.font.name = 'Calibri'
    r1 = row.cells[1].paragraphs[0].add_run(usage)
    r1.font.size = Pt(10)
    r1.font.name = 'Calibri'
    if i % 2 == 0:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'F9F9F9')
            tc_pr.append(shd)

doc.add_paragraph()

# ════════════════════════════════════════════════════════════════
#  HTML CODE SAMPLE
# ════════════════════════════════════════════════════════════════
section_heading(doc, 'Code Sample — HTML (Hero Section)')

html_code = '''<!-- HERO SECTION -->
<section id="home" class="hero-section">
  <div class="hero-overlay">
    <div class="container">
      <div class="row align-items-center min-vh-100 py-5">
        <div class="col-lg-7">
          <h1 class="hero-title">
            Helping Business And Companies
            Innovate Transform And Lead
          </h1>
          <p class="hero-text">
            We partner with businesses of all sizes to create
            meaningful digital experiences.
          </p>
          <div class="hero-buttons">
            <a href="#contact" class="btn-red-solid">Start Now</a>
            <a href="#services" class="btn-white-outline">Our Services</a>
          </div>
        </div>
      </div>
    </div>
  </div>
</section>'''

p = doc.add_paragraph()
shade_paragraph(p, 'F4F4F4')
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after  = Pt(4)
p.paragraph_format.left_indent  = Cm(0.4)
r = p.add_run(html_code)
r.font.name = 'Courier New'
r.font.size  = Pt(8)
r.font.color.rgb = RGBColor(0x1A, 0x1E, 0x2A)

doc.add_paragraph()

# ════════════════════════════════════════════════════════════════
#  CSS CODE SAMPLE
# ════════════════════════════════════════════════════════════════
section_heading(doc, 'Code Sample — CSS (Hero & Feature Cards)')

css_code = ''':root {
  --red:      #e74c3c;
  --dark:     #1a1e2a;
  --dark-card:#2c3038;
  --font:     'Poppins', sans-serif;
}

html { scroll-behavior: smooth; }

/* Hero Section */
.hero-section {
  background-image: url('https://images.unsplash.com/...');
  background-size: cover;
  background-position: center;
  min-height: 100vh;
}

.hero-overlay {
  background: rgba(15, 18, 28, 0.82);
  min-height: 100vh;
  display: flex;
  align-items: center;
  padding-top: 80px;
}

.hero-title {
  font-size: clamp(2rem, 4vw, 3.2rem);
  font-weight: 800;
  color: #ffffff;
  line-height: 1.25;
  margin-bottom: 22px;
}

/* Feature Cards */
.feature-cards-section {
  position: relative;
  z-index: 10;
  margin-top: -80px;   /* overlaps hero bottom */
}

.feature-card { padding: 44px 36px; color: #fff; }
.card-red     { background-color: var(--red); }
.card-dark    { background-color: var(--dark-card); }
.feature-card:hover { transform: translateY(-6px); }

/* Responsive */
@media (max-width: 767px) {
  .hero-title { font-size: 1.8rem; }
  .feature-cards-section { margin-top: 0; }
}'''

p2 = doc.add_paragraph()
shade_paragraph(p2, 'F4F4F4')
p2.paragraph_format.space_before = Pt(4)
p2.paragraph_format.space_after  = Pt(4)
p2.paragraph_format.left_indent  = Cm(0.4)
r2 = p2.add_run(css_code)
r2.font.name = 'Courier New'
r2.font.size  = Pt(8)
r2.font.color.rgb = RGBColor(0x1A, 0x1E, 0x2A)

doc.add_paragraph()

# ════════════════════════════════════════════════════════════════
#  SCREENSHOT PLACEHOLDER NOTE
# ════════════════════════════════════════════════════════════════
section_heading(doc, 'Web Page Screenshots')

note = doc.add_paragraph()
shade_paragraph(note, 'FFF3CD')
note.paragraph_format.space_before = Pt(6)
note.paragraph_format.space_after  = Pt(6)
note.paragraph_format.left_indent  = Cm(0.4)
nr = note.add_run(
    'Screenshots of the live web page are attached below. '
    'The live site can be viewed at:\n'
    'https://stirring-snickerdoodle-5a985f.netlify.app\n\n'
    'Sections visible in the screenshots:\n'
    '  • Navigation Bar & Hero Section\n'
    '  • Feature Cards (red, dark, darker)\n'
    '  • Services Section (6 services)\n'
    '  • Dark CTA Section\n'
    '  • Why Us / About Section\n'
    '  • Portfolio Gallery (8 images)\n'
    '  • Stats, Clients, Features Sections\n'
    '  • Contact Form\n'
    '  • Footer'
)
nr.font.size = Pt(10)
nr.font.name = 'Calibri'
nr.font.color.rgb = RGBColor(0x85, 0x64, 0x04)

# ════════════════════════════════════════════════════════════════
#  FOOTER
# ════════════════════════════════════════════════════════════════
doc.add_paragraph()
add_hr(doc, '1A1E2A')
footer_p = doc.add_paragraph()
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.paragraph_format.space_before = Pt(8)
fr = footer_p.add_run('Fazilma Khan Niazi  |  Roll No: S2024266184  |  Assignment 2  |  Web Development')
fr.font.size = Pt(9)
fr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
fr.font.name = 'Calibri'

# ── Save ──────────────────────────────────────────────────────
doc.save('/home/user/coursera-test/assignment2/Assignment2_FazilmaKhanNiazi_S2024266184.docx')
print("Word document created successfully!")
